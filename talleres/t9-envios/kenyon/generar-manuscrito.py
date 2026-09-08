import json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

paras = json.load(open('paras.json'))
d = Document()

# --- estilo Normal: Times New Roman 12, doble espacio, alineado a la izquierda
st = d.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
rpr = st.element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
    rf.set(qn(a), 'Times New Roman')
pf = st.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
pf.space_before = Pt(0); pf.space_after = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m in ('top_margin','bottom_margin','left_margin','right_margin'):
    setattr(sec, m, Inches(1))

# --- encabezado: Gamboa / THE RECEIPT / nº
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run('Gamboa / THE RECEIPT / ')
r = hp.add_run()
for el, attr in (('w:fldChar', ('w:fldCharType', 'begin')), None, ('w:fldChar', ('w:fldCharType', 'end'))) if False else ():
    pass
def field(run, instr):
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
field(r, ' PAGE ')

def P(text, center=False, indent=False, single=False):
    p = d.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    f = p.paragraph_format
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE if single else WD_LINE_SPACING.DOUBLE
    f.space_before = Pt(0); f.space_after = Pt(0)
    f.first_line_indent = Inches(0.5) if indent else Inches(0)
    return p

for line in ['Dennis Gamboa', 'Madrid, Spain', 'private.hera@proton.me', 'Approx. 1,300 words']:
    P(line, single=True)
for _ in range(4): P('')
P('THE RECEIPT', center=True)
P('')
for t in paras:
    if t.startswith('There is a plastic bag'):
        P('#', center=True)
    P(t, indent=True)
P('END', center=True)

d.core_properties.author = 'Dennis Gamboa'
d.core_properties.last_modified_by = 'Dennis Gamboa'
d.core_properties.title = 'The Receipt'
d.save('The Receipt - Dennis Gamboa.docx')
print('ok')
